from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from playwright.sync_api import sync_playwright

BASE = Path(__file__).parent.resolve()
OUT = BASE / "心理学知识体系_图文版.docx"
SHOT_DIR = BASE / "word_screenshots"
SHOT_DIR.mkdir(exist_ok=True)


def file_url(name: str) -> str:
    return (BASE / name).resolve().as_uri()


def compress_png(path: Path, max_width=2200):
    img = Image.open(path)
    if img.width > max_width:
        ratio = max_width / img.width
        img = img.resize((max_width, int(img.height * ratio)))
        img.save(path)


def capture_screenshots():
    shots = {}
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 2400, "height": 1700}, device_scale_factor=1)

        page.goto(file_url("knowledge_top.html"), wait_until="networkidle")
        page.wait_for_selector("svg", timeout=30000)
        page.wait_for_timeout(1200)
        path = SHOT_DIR / "01_top_mindmap.png"
        page.screenshot(path=str(path), full_page=True)
        compress_png(path)
        shots["top"] = path

        page.goto(file_url("knowledge_middle.html"), wait_until="networkidle")
        page.wait_for_selector("#filterVal", timeout=30000)
        for label in ["入门", "进阶", "高级"]:
            page.select_option("#filterVal", label=label)
            page.wait_for_timeout(500)
            path = SHOT_DIR / f"02_middle_{label}.png"
            page.screenshot(path=str(path), full_page=True)
            compress_png(path)
            shots[f"middle_{label}"] = path

        page.goto(file_url("knowledge_bottom.html"), wait_until="networkidle")
        page.wait_for_selector("canvas", timeout=30000)
        page.wait_for_timeout(1400)
        path = SHOT_DIR / "03_bottom_graph.png"
        page.screenshot(path=str(path), full_page=True)
        compress_png(path)
        shots["bottom"] = path

        page.goto(file_url("knowledge_learning_path.html"), wait_until="networkidle")
        page.wait_for_selector("g.node", timeout=30000)
        page.wait_for_timeout(1000)
        routes = [
            ("全景路线", "all"),
            ("通识入门", "intro"),
            ("认知路线", "cognition"),
            ("发展人格", "development_personality"),
            ("社会路线", "social"),
            ("临床健康", "clinical_health"),
            ("前沿路线", "frontier"),
        ]
        for label, value in routes:
            page.select_option("#route", value=value)
            page.wait_for_timeout(900)
            path = SHOT_DIR / f"04_route_{value}.png"
            page.screenshot(path=str(path), full_page=False)
            compress_png(path)
            shots[f"route_{value}"] = path

        browser.close()
    return shots


def set_document_style(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)


def add_page_break_before(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:pageBreakBefore"))


def add_markdown(doc: Document, md_path: Path):
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.style = doc.styles["Intense Quote"]
        else:
            doc.add_paragraph(line)


def add_picture_block(doc: Document, title: str, image_path: Path, description: str):
    doc.add_heading(title, level=2)
    doc.add_paragraph(description)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.8))


def add_visual_sections(doc: Document, shots):
    h = doc.add_heading("四、可视化建模截图与解读", level=1)
    add_page_break_before(h)
    doc.add_paragraph("以下截图由 Playwright Chromium 无头模式自动打开本地 HTML 页面后生成，与心理学知识体系四个可视化文件一一对应。")

    add_picture_block(
        doc,
        "4.1 顶层思维导图：心理学主干",
        shots["top"],
        "顶层图展示心理学的十个一级主干：心理学导论与方法、生物与神经基础、感觉知觉与注意、学习记忆与认知、发展心理学、人格与个体差异、社会心理学、异常与临床心理、健康与咨询应用、认知神经与前沿。"
    )

    doc.add_heading("4.2 中间层矩阵：按难度层级截图", level=2)
    descriptions = {
        "入门": "入门矩阵聚焦科学心理学的基础框架，包括心理学科学观、实验设计、相关与因果、神经基础、学习机制、发展、人 格、社会认知、心理障碍分类和 CBT 基础。",
        "进阶": "进阶矩阵聚焦测量、机制和应用边界，包括信效度、研究伦理、注意与执行控制、记忆系统、情绪调节、社会影响、临床评估和健康行为改变。",
        "高级": "高级矩阵聚焦复杂情境、临床风险和前沿方法，包括危机干预、脑成像与认知神经、开放科学与可重复性。"
    }
    for level in ["入门", "进阶", "高级"]:
        add_picture_block(doc, f"4.2 {level}矩阵", shots[f"middle_{level}"], descriptions[level])

    add_picture_block(
        doc,
        "4.3 最底层知识网络图谱",
        shots["bottom"],
        "底层图谱展示 72 个心理学核心知识节点与 76 条关系边。节点大小由重要度和枢纽度共同决定，关系类型包括 depends、contains、evolved、analogy、contrast。"
    )

    doc.add_heading("4.4 学习路线图：按路线截图", level=2)
    route_desc = {
        "all": ("全景路线", "展示完整知识网络，适合建立整体结构和复盘覆盖范围。"),
        "intro": ("通识入门", "从科学方法和行为开始，进入生物基础、感觉知觉、学习记忆、发展、人格、社会和临床。"),
        "cognition": ("认知路线", "从神经基础进入感觉、注意、工作记忆、长时记忆、问题解决和认知模型。"),
        "development_personality": ("发展人格", "关注生命周期发展、依恋、认知发展、身份认同、人格特质和情绪调节。"),
        "social": ("社会路线", "围绕态度、归因、偏见、从众、服从、说服和群体决策展开。"),
        "clinical_health": ("临床健康", "从压力进入焦虑、抑郁、创伤、访谈、量表、CBT、压力管理和危机干预。"),
        "frontier": ("前沿路线", "连接科学方法、研究伦理、脑成像、认知模型、机器学习心理学、人机交互和开放科学。"),
    }
    for key in ["all", "intro", "cognition", "development_personality", "social", "clinical_health", "frontier"]:
        title, desc = route_desc[key]
        add_picture_block(doc, f"4.4 {title}", shots[f"route_{key}"], desc)


def main():
    shots = capture_screenshots()
    doc = Document()
    set_document_style(doc)
    add_markdown(doc, BASE / "知识体系_心理学.md")
    add_visual_sections(doc, shots)
    doc.add_page_break()
    doc.add_heading("附录：截图文件清单", level=1)
    for name, path in shots.items():
        doc.add_paragraph(f"{name}: {path.name}", style="List Bullet")
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
